import warnings
warnings.filterwarnings("ignore", category=UserWarning)
import cv2
import easyocr
import pandas as pd
from ultralytics import YOLO
from datetime import datetime
import os
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

# --- 1. CONFIGURATION & KNOWLEDGE BASE ---
LOG_FILE = "video_log.xlsx"
VIDEO_PATH = "sample_video.mp4" 

NIGERIA_STATES = {
    "ABJ": "Abuja (FCT)", "FCT": "Abuja (FCT)", "ABIA": "Umuahia", "ADAMAWA": "Yola", 
    "AKWA IBOM": "Uyo", "ANAMBRA": "Awka", "ANA": "Awka", "BAUCHI": "Bauchi", 
    "BAYELSA": "Yenagoa", "BENUE": "Makurdi", "BEN": "Makurdi", "BORNO": "Maiduguri", 
    "CROSS RIVER": "Calabar", "DELTA": "Asaba", "DEL": "Asaba", "EBONYI": "Abakaliki", 
    "EDO": "Benin City", "EKITI": "Ado-Ekiti", "ENUGU": "Enugu", "ENU": "Enugu", 
    "GOMBE": "Gombe", "IMO": "Owerri", "JIGAWA": "Dutse", "KADUNA": "Kaduna", 
    "KAD": "Kaduna", "KANO": "Kano", "KAN": "Kano", "KATSINA": "Katsina", 
    "KEBBI": "Birnin Kebbi", "KOGI": "Lokoja", "KWARA": "Ilorin", "LAGOS": "Ikeja", 
    "LAG": "Ikeja", "NASARAWA": "Lafia", "NIGER": "Minna", "OGUN": "Abeokuta", 
    "OGU": "Abeokuta", "ONDO": "Akure", "OND": "Akure", "OSUN": "Osogbo", 
    "OYO": "Ibadan", "PLATEAU": "Jos", "RIVERS": "Port Harcourt", 
    "PHC": "Port Harcourt", "SOKOTO": "Sokoto", "TARABA": "Jalingo", 
    "YOBE": "Damaturu", "ZAMFARA": "Gusau"
}

def identify_origin(ocr_text, plate_number):
    text = ocr_text.upper().replace(" ", "").replace("-", "")
    for keyword, capital in NIGERIA_STATES.items():
        if keyword in text: return f"Nigeria: {keyword.title()} ({capital})"
    nigeria_pattern = r'^[A-Z]{3}\d{3}[A-Z]{2}$|^[A-Z]{2}\d{3}[A-Z]{3}$'
    if re.match(nigeria_pattern, plate_number): return "Nigeria (General)"
    return "International"

# --- 2. EXCEL STYLING (Blue Header + Auto-Width) ---
def format_excel_report(file_path):
    try:
        wb = load_workbook(file_path)
        ws = wb.active
        header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        
        for cell in ws[1]: # Styles all headers including Time_Stamp
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            
        for col in ws.columns:
            max_length = 0
            column_letter = col[0].column_letter
            for cell in col:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            ws.column_dimensions[column_letter].width = max_length + 5
        wb.save(file_path)
    except Exception as e: print(f"Excel Error: {e}")

# --- 3. INITIALIZATION & FITTED WINDOW ---
model = YOLO("yolo11n.pt")
reader = easyocr.Reader(['en'], gpu=False)
cap = cv2.VideoCapture(VIDEO_PATH)
seen_plates = set()

# Window Setup (Fitted & Top-Left)
target_w, target_h = 854, 480
WINDOW_NAME = "ALPR Video Analysis (Fitted)"
cv2.namedWindow(WINDOW_NAME, cv2.WINDOW_NORMAL)
cv2.resizeWindow(WINDOW_NAME, target_w, target_h)
cv2.moveWindow(WINDOW_NAME, 0, 0)

headers = ["Time_Stamp", "Model", "Origin_Capital", "Plate_Number", "Confidence"]
if not os.path.exists(LOG_FILE):
    pd.DataFrame(columns=headers).to_excel(LOG_FILE, index=False)
    format_excel_report(LOG_FILE)

# --- 4. MAIN LOOP ---
while cap.isOpened():
    ret, frame = cap.read()
    if not ret: break

    results = model(frame, conf=0.20, verbose=False)
    for result in results:
        for box in result.boxes:
            x1, y1, x2, y2 = map(int, box.xyxy[0])
            pad = 20 # Resolves 'Unknown' state
            y1_p = max(0, y1 - pad)
            plate_crop = frame[y1_p:y2, x1:x2]

            if plate_crop.size > 0:
                plate_res = cv2.resize(plate_crop, None, fx=2.5, fy=2.5, interpolation=cv2.INTER_CUBIC)
                ocr_res = reader.readtext(plate_res)

                for (_, text, prob) in ocr_res:
                    clean_num = re.sub(r'[^A-Z0-9]', '', text.upper())
                    if len(clean_num) >= 5 and clean_num not in seen_plates and prob > 0.45:
                        origin = identify_origin(text, clean_num)
                        new_data = pd.DataFrame([{"Time_Stamp": datetime.now().strftime("%H:%M:%S"), "Model": "YOLOv11n (Video)", "Origin_Capital": origin, "Plate_Number": clean_num, "Confidence": round(prob, 2)}])
                        with pd.ExcelWriter(LOG_FILE, mode='a', engine='openpyxl', if_sheet_exists='overlay') as writer:
                            new_data.to_excel(writer, index=False, header=False, startrow=writer.sheets['Sheet1'].max_row)
                        format_excel_report(LOG_FILE)
                        seen_plates.add(clean_num)
                        cv2.rectangle(frame, (x1, y1), (x2, y2), (0, 255, 0), 2)

    # Sharp fitted display
    display_frame = cv2.resize(frame, (target_w, target_h), interpolation=cv2.INTER_AREA)
    cv2.imshow(WINDOW_NAME, display_frame)
    if cv2.waitKey(1) & 0xFF == ord('q'): break

cap.release()
cv2.destroyAllWindows()


