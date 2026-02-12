import cv2
import easyocr
import pandas as pd
from ultralytics import YOLO
from datetime import datetime
import os
import re
import difflib
from openpyxl.styles import PatternFill, Font, Alignment

# --- CONFIGURATION ---
INPUT_FOLDER = "test_images"       
OUTPUT_FOLDER = "test_results"     
CONF_THRESHOLD = 0.20              

os.makedirs(OUTPUT_FOLDER, exist_ok=True)

# --- 1. KNOWLEDGE BASE ---
LGA_MAP = {
    "KRD": "Lagos", "APP": "Lagos", "LND": "Lagos", "MUS": "Lagos",
    "EKY": "Lagos", "FKJ": "Lagos", "GGE": "Lagos", "LSR": "Lagos",
    "KTU": "Lagos", "RSH": "Abuja (FCT)", "ABC": "Abuja (FCT)",
    "KUJ": "Abuja (FCT)", "BWR": "Abuja (FCT)", "MKA": "Kaduna",
    "DKA": "Kaduna", "KAN": "Kano", "PHC": "Rivers"
}

NIGERIA_STATES = {
    "ABJ": "Abuja (FCT)", "ABUJA": "Abuja (FCT)", "FCT": "Abuja (FCT)",
    "BEN": "Benue", "BENUE": "Benue", "KOG": "Kogi", "KOGI": "Kogi",
    "KWA": "Kwara", "KWARA": "Kwara", "NAS": "Nasarawa", "NASARAWA": "Nasarawa",
    "NIG": "Niger", "NIGER": "Niger", "PLA": "Plateau", "PLATEAU": "Plateau", "JOS": "Plateau",
    "ADA": "Adamawa", "ADAMAWA": "Adamawa", "BAU": "Bauchi", "BAUCHI": "Bauchi",
    "BOR": "Borno", "BORNO": "Borno", "GOM": "Gombe", "GOMBE": "Gombe",
    "TAR": "Taraba", "TARABA": "Taraba", "YOB": "Yobe", "YOBE": "Yobe",
    "JIG": "Jigawa", "JIGAWA": "Jigawa", "KAD": "Kaduna", "KADUNA": "Kaduna",
    "KAN": "Kano", "KANO": "Kano", "KAT": "Katsina", "KATSINA": "Katsina",
    "KEB": "Kebbi", "KEBBI": "Kebbi", "SOK": "Sokoto", "SOKOTO": "Sokoto",
    "ZAM": "Zamfara", "ZAMFARA": "Zamfara", "ABI": "Abia", "ABIA": "Abia",
    "ANA": "Anambra", "ANAMBRA": "Anambra", "EBO": "Ebonyi", "EBONYI": "Ebonyi",
    "ENU": "Enugu", "ENUGU": "Enugu", "IMO": "Imo", "AKW": "Akwa Ibom", "AKWA": "Akwa Ibom",
    "IBOM": "Akwa Ibom", "BAY": "Bayelsa", "BAYELSA": "Bayelsa", "CRO": "Cross River",
    "CROSS": "Cross River", "DEL": "Delta", "DELTA": "Delta", "EDO": "Edo",
    "RIV": "Rivers", "RIVERS": "Rivers", "PHC": "Rivers", "PH": "Rivers",
    "EKI": "Ekiti", "EKITI": "Ekiti", "LAG": "Lagos", "LAGOS": "Lagos",
    "CENTRE": "Lagos", "LND": "Lagos", "OGU": "Ogun", "OGUN": "Ogun",
    "OND": "Ondo", "ONDO": "Ondo", "SUN": "Ondo", "OSU": "Osun", "OSUN": "Osun", "OYO": "Oyo"
}

def identify_state(text, plate_number=""):
    clean_text = text.upper().replace(" ", "").replace("-", "")
    for keyword, full_name in NIGERIA_STATES.items():
        if keyword in clean_text: return full_name
    if len(plate_number) > 3:
        prefix = plate_number[:3].upper()
        if prefix in LGA_MAP: return LGA_MAP[prefix]
    matches = difflib.get_close_matches(clean_text, list(NIGERIA_STATES.keys()), n=1, cutoff=0.8)
    if matches: return NIGERIA_STATES[matches[0]]
    return "Unknown"

# --- 2. CLEANERS ---
def clean_plate_text(raw_text):
    clean = re.sub(r'[^A-Z0-9]', '', raw_text.upper())
    if len(clean) > 10: return "Unreadable (Noise)"
    if len(clean) < 6: return clean 
    chars = list(clean)
    def fix_num(c): return {'O': '0', 'I': '1', 'S': '5', 'B': '8', 'Z': '2', 'A': '4', 'G': '6'}.get(c, c)
    def fix_let(c): return {'0': 'O', '1': 'I', '5': 'S', '8': 'B', '4': 'A', '6': 'G'}.get(c, c)
    if len(chars) == 8:
        chars[0], chars[1], chars[2] = fix_let(chars[0]), fix_let(chars[1]), fix_let(chars[2])
        chars[3], chars[4], chars[5] = fix_num(chars[3]), fix_num(chars[4]), fix_num(chars[5])
        chars[6], chars[7] = fix_let(chars[6]), fix_let(chars[7])
        return f"{''.join(chars[:3])}-{''.join(chars[3:6])}-{''.join(chars[6:])}"
    return "".join(chars)

def clean_general_plate(raw_text):
    text = raw_text.upper()
    SLOGANS = ["FEDERAL REPUBLIC OF NIGERIA", "FEDERAL", "REPUBLIC", "NIGERIA", "CENTRE OF EXCELLENCE", "CENTER OF EXCELLENCE", "HEARTBEAT OF THE NATION", "THE STATE OF HARMONY", "COAL CITY STATE", "SALT OF THE NATION", "LAGOS", "ABUJA", "FCT", "PEACE AND TOURISM"]
    for slogan in SLOGANS: text = text.replace(slogan, "")
    return re.sub(r'[^A-Z0-9-]', '', text)

# --- 3. LOAD MODELS ---
print("1. Loading Models...")
model_ng = YOLO("best.pt")
reader = easyocr.Reader(['en'], gpu=False)

log_file = "batch_results_log.xlsx"
df_log = pd.DataFrame(columns=["Time_Stamp", "Image_Name", "Model", "Origin_Capital", "Plate_Number", "Confidence"])

# --- 4. PROCESS FUNCTION ---
def process_batch():
    global df_log
    
    # Create evidence folder
    evidence_folder = os.path.join(OUTPUT_FOLDER, "evidence_plates")
    os.makedirs(evidence_folder, exist_ok=True)

    files = [f for f in os.listdir(INPUT_FOLDER) if f.lower().endswith(('.jpg', '.jpeg', '.png'))]
    print(f"2. Found {len(files)} images. Processing...")

    for filename in files:
        img_path = os.path.join(INPUT_FOLDER, filename)
        img = cv2.imread(img_path)
        if img is None: continue

        print(f"   -> Processing: {filename}...")
        current_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        # TEST 1: NIGERIA
        results_ng = model_ng(img, verbose=False)
        for result in results_ng:
            for box in result.boxes:
                score = float(box.conf[0])
                if score > CONF_THRESHOLD:
                    x1, y1, x2, y2 = map(int, box.xyxy[0])
                    plate_crop = img[y1:y2, x1:x2]
                    
                    # SAVE EVIDENCE
                    gray = cv2.cvtColor(plate_crop, cv2.COLOR_BGR2GRAY)
                    contrast = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
                    ocr_res = reader.readtext(contrast)
                    detected_state = "Unknown"
                    potential_numbers = []
                    for (_, text, prob) in ocr_res:
                        if prob > 0.1 and len(text) > 2:
                            found_state = identify_state(text)
                            if found_state != "Unknown": detected_state = found_state 
                            else: potential_numbers.append(clean_plate_text(text))
                    detected_number = max(potential_numbers, key=len) if potential_numbers else "Unreadable"
                    if detected_state == "Unknown": detected_state = identify_state("", detected_number)
                    
                    safe_name = re.sub(r'[^A-Z0-9]', '', detected_number)
                    evidence_path = os.path.join(evidence_folder, f"EVIDENCE_{safe_name}_{filename}")
                    cv2.imwrite(evidence_path, plate_crop)

                    cv2.rectangle(img, (x1, y1), (x2, y2), (0, 255, 0), 2)
                    cv2.putText(img, f"{detected_number}", (x1, y1 - 10), cv2.FONT_HERSHEY_SIMPLEX, 0.6, (0, 255, 0), 2)
                    new_entry = {"Time_Stamp": current_time, "Image_Name": filename, "Model": "Nigeria", "Origin_Capital": detected_state, "Plate_Number": detected_number, "Confidence": round(score, 2)}
                    df_log = pd.concat([df_log, pd.DataFrame([new_entry])], ignore_index=True)

       
    # --- 5. COLORFUL EXCEL SAVING (WITH HEIGHT & ALIGNMENT) ---
    with pd.ExcelWriter(log_file, engine='openpyxl') as writer:
        df_log.to_excel(writer, index=False, sheet_name='Sheet1')
        ws = writer.sheets['Sheet1']
        
        # Styles
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark Blue
        header_font = Font(color="FFFFFF", bold=True)
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        
        # --- NEW: Set Header Height to 22 ---
        ws.row_dimensions[1].height = 22

        # Apply to Header
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            # --- NEW: Middle and Center Align ---
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Apply to Data Rows
        for row in range(2, ws.max_row + 1):
            model_cell = ws[f'C{row}']
            if model_cell.value == "Nigeria":
                for col in range(1, 7): ws.cell(row=row, column=col).fill = green_fill
            elif model_cell.value == "International":
                for col in range(1, 7): ws.cell(row=row, column=col).fill = red_fill

        # Set Column Widths to 25
        for col in ['A', 'B', 'C', 'D', 'E', 'F']:
            ws.column_dimensions[col].width = 25
            
    print(f"\n3. DONE! Colorful formatted log saved to '{log_file}'. Evidence saved.")

process_batch()