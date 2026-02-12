import warnings
warnings.filterwarnings("ignore", category=UserWarning)
import cv2
import easyocr
import pandas as pd
from ultralytics import YOLO
from datetime import datetime
import os
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

# --- 1. CONFIGURATION ---
LOG_FILE = "nigeria_alpr_cpu.xlsx"

# Standardizing device to CPU to avoid CUDA conflicts
DEVICE = 'cpu'

NIGERIA_STATES = {
    "ABJ": "Abuja (FCT)", "FCT": "Abuja (FCT)", "ABIA": "Umuahia", "ADAMAWA": "Yola", 
    "AKWA IBOM": "Uyo", "ANAMBRA": "Awka", "ANA": "Awka", "BAUCHI": "Bauchi", 
    "BAYELSA": "Yenagoa", "BENUE": "Makurdi", "BEN": "Makurdi", "BORNO": "Maiduguri", 
    "CROSS RIVER": "Calabar", "DELTA": "Asaba", "DEL": "Asaba", "EBONYI": "Abakaliki", 
    "EDO": "Benin City", "EKITI": "Ado-Ekiti", "ENUGU": "Enugu", "ENU": "Enugu", 
    "GOMBE": "Gombe", "IMO": "Owerri", "JIGAWA": "Dutse", "KADUNA": "Kaduna", 
    "KAD": "Kaduna", "KANO": "Kano", "KAN": "Kano", "KATSINA": "Katsina", 
    "KEBBI": "Birnin Kebbi", "KOGI": "Lokoja", "KWARA": "Ilorin", "LAGOS": "Ikeja", 
    "LAG": "Ikeja", "NASARAWA": "Lafia", "NIGER": "Minna", "OGUN": "Abeokuta", 
    "OGU": "Abeokuta", "ONDO": "Akure", "OND": "Akure", "OSUN": "Osogbo", 
    "OYO": "Ibadan", "PLATEAU": "Jos", "RIVERS": "Port Harcourt", 
    "PHC": "Port Harcourt", "SOKOTO": "Sokoto", "TARABA": "Jalingo", 
    "YOBE": "Damaturu", "ZAMFARA": "Gusau"
}

# --- 2. INITIALIZATION ---
# Loading your custom weights and forcing them to CPU
model = YOLO("best.pt").to(DEVICE)

# Disabling GPU for EasyOCR
reader = easyocr.Reader(['en'], gpu=False)

# Camera setup with reduced resolution for speed
cap = cv2.VideoCapture(0, cv2.CAP_DSHOW)
cap.set(cv2.CAP_PROP_FRAME_WIDTH, 640)
cap.set(cv2.CAP_PROP_FRAME_HEIGHT, 480)
cap.set(cv2.CAP_PROP_BUFFERSIZE, 1)

seen_plates = set()
frame_count = 0

# --- 3. CORE LOGIC ---
def identify_origin(ocr_text, plate_number):
    text = ocr_text.upper().replace(" ", "").replace("-", "")
    for keyword, capital in NIGERIA_STATES.items():
        if keyword in text: 
            return f"Nigeria: {keyword.title()} ({capital})"
    
    nigeria_pattern = r'^[A-Z]{3}\d{3}[A-Z]{2}$|^[A-Z]{2}\d{3}[A-Z]{3}$'
    if re.match(nigeria_pattern, plate_number): 
        return "Nigeria (General)"
    return "Unknown/International"

while cap.isOpened():
    ret, frame = cap.read()
    if not ret: break

    frame_count += 1
    # Optimization: Only perform heavy AI processing every 3rd frame to save CPU
    if frame_count % 3 != 0:
        cv2.imshow("ALPR Live Monitoring (CPU Optimized)", frame)
        if cv2.waitKey(1) & 0xFF == ord('q'): break
        continue

    # Detection
    results = model(frame, conf=0.15
                    , verbose=False, device=DEVICE)

    for result in results:
        for box in result.boxes:
            x1, y1, x2, y2 = map(int, box.xyxy[0])
            
            # Padding to ensure "State" name is included in crop
            pad = 20
            crop = frame[max(0, y1-pad):y2, x1:x2]

            if crop.size > 0:
                # OCR processing
                ocr_results = reader.readtext(crop)
                for (_, text, prob) in ocr_results:
                    clean_num = re.sub(r'[^A-Z0-9]', '', text.upper())
                    
                    if len(clean_num) >= 7 and clean_num not in seen_plates:
                        origin = identify_origin(text, clean_num)
                       
                        # Data Logging
                        new_log = pd.DataFrame([{
                            "Timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                            "Plate": clean_num,
                            "Origin": origin,
                            "Confidence": round(prob, 2)
                        }])
                        
                        if not os.path.exists(LOG_FILE):
                            new_log.to_excel(LOG_FILE, index=False)
                        else:
                            with pd.ExcelWriter(LOG_FILE, mode='a', engine='openpyxl', if_sheet_exists='overlay') as writer:
                                new_log.to_excel(writer, index=False, header=False, startrow=writer.sheets['Sheet1'].max_row)
                        
                        seen_plates.add(clean_num)
                        cv2.rectangle(frame, (x1, y1), (x2, y2), (0, 255, 0), 2)

    cv2.imshow("ALPR Live Monitoring (CPU Optimized)", frame)
    if cv2.waitKey(1) & 0xFF == ord('q'): break

cap.release()
cv2.destroyAllWindows()


